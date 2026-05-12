from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

from app.config import PILOTS_DIR
from app.master_data_template import (
    META_SHEETS,
    SHEET_COLUMNS,
    _allowed_values_df,
    _autosize_columns,
    _data_dictionary_df,
    _instructions_df,
)


PILOT_PATH = PILOTS_DIR / "maria_matias_april_2026_master_data.xlsx"

# Pilot source note:
# This workbook is reconstructed from a submitted April 2026 claim example.
# It is not an official approved timetable and must not be treated as one.
# The original ICT Distance row was described as "Test" on the claim.
# Phase 1.3.1 intentionally records it as a claimable teaching session
# treated as Lecture because the tariff and verification rules are the same.

EXPECTED_SESSIONS = 69
EXPECTED_HOURS = 94.00
EXPECTED_AMOUNT = 43240.00


TIMETABLE_ROWS = [
    ["1008977", "CUS HORTICULTURE", "CUS411S", "Wednesday", "10:30", "12:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "CUS HORTICULTURE", "CUS411S", "Thursday", "08:30", "10:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "CUS GROUP A", "CUS411S", "Monday", "08:30", "09:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "CUS GROUP A", "CUS411S", "Tuesday", "11:30", "12:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "CUS GROUP A", "CUS411S", "Wednesday", "07:30", "08:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "CUS GROUP A", "CUS411S", "Thursday", "14:00", "15:00", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "CUS GROUP B", "CUS411S", "Monday", "07:30", "08:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "CUS GROUP B", "CUS411S", "Tuesday", "10:30", "11:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "CUS GROUP B", "CUS411S", "Thursday", "15:00", "16:00", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "CUS GROUP B", "CUS411S", "Friday", "07:30", "08:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "ICT GROUP A", "ICT521S", "Monday", "11:30", "13:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "ICT GROUP A", "ICT521S", "Tuesday", "08:30", "10:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "ICT GROUP B", "ICT521S", "Monday", "09:30", "11:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "ICT GROUP B", "ICT521S", "Tuesday", "14:00", "16:00", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "ICT BOA-EENANHA", "ICT521S", "Thursday", "10:30", "12:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "ICT BOA-EENANHA", "ICT521S", "Friday", "11:30", "13:00", "2026-04-10", "2026-04-10", "Yes"],
    ["1008977", "ICT BOA-EENANHA", "ICT521S", "Friday", "09:30", "10:30", "2026-04-17", "2026-04-17", "Yes"],
    ["1008977", "ICT BOA-EENANHA", "ICT521S", "Friday", "08:30", "10:00", "2026-04-24", "2026-04-24", "Yes"],
    ["1008977", "ICT GREY", "ICT521S", "Monday", "14:00", "15:00", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "ICT GREY", "ICT521S", "Tuesday", "07:30", "08:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "ICT GREY", "ICT521S", "Wednesday", "08:30", "09:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "ICT GREY", "ICT521S", "Friday", "10:30", "11:30", "2026-04-07", "2026-04-29", "Yes"],
    ["1008977", "ICT Distance", "ICT521S", "Saturday", "09:00", "10:00", "2026-04-18", "2026-04-18", "Yes"],
]


CALENDAR_ROWS = [
    ["New Year's Day", "2026-01-01", "2026-01-01", "public_holiday", "exclude", "No"],
    ["Semester 1 Mid-Semester Break", "2026-03-30", "2026-04-02", "recess", "exclude", "No"],
    ["Independence Day", "2026-03-21", "2026-03-21", "public_holiday", "exclude", "No"],
    ["Good Friday", "2026-04-03", "2026-04-03", "public_holiday", "exclude", "No"],
    ["Easter Sunday", "2026-04-05", "2026-04-05", "public_holiday", "exclude", "No"],
    ["Easter Monday", "2026-04-06", "2026-04-06", "public_holiday", "exclude", "No"],
    ["Workers' Day", "2026-05-01", "2026-05-01", "public_holiday", "exclude", "No"],
    ["Cassinga Day", "2026-05-04", "2026-05-04", "public_holiday", "exclude", "No"],
    ["Ascension Day", "2026-05-14", "2026-05-14", "public_holiday", "exclude", "No"],
    ["Africa Day", "2026-05-25", "2026-05-25", "public_holiday", "exclude", "No"],
    ["Genocide Remembrance Day", "2026-05-28", "2026-05-28", "public_holiday", "exclude", "No"],
    ["Institutional Holiday", "2026-05-29", "2026-05-29", "institutional_closure", "exclude", "No"],
    ["Mid-Year Recess - Students", "2026-06-15", "2026-07-10", "recess", "exclude", "No"],
    ["Semester 2 Mid-Semester Break", "2026-09-07", "2026-09-11", "recess", "exclude", "No"],
    ["Heroes' Day", "2026-08-26", "2026-08-26", "public_holiday", "exclude", "No"],
    ["Day of the Namibian Women and International Human Rights Day", "2026-12-10", "2026-12-10", "public_holiday", "exclude", "No"],
    ["Christmas Day", "2026-12-25", "2026-12-25", "public_holiday", "exclude", "No"],
    ["Family Day", "2026-12-26", "2026-12-26", "public_holiday", "exclude", "No"],
]


def _pilot_instructions_df() -> pd.DataFrame:
    pilot_rows = pd.DataFrame(
        [
            ["Pilot notice", "Maria Matias April 2026 pilot reconstructed from a submitted claim example."],
            ["Pilot warning", "This is not the official approved timetable. Use it only to test workflow behavior."],
            ["Privacy", "Sensitive lecturer fields use placeholders. Student rows are dummy-only."],
        ],
        columns=["topic", "instruction"],
    )
    return pd.concat([pilot_rows, _instructions_df()], ignore_index=True)


def _pilot_dataframes() -> dict[str, pd.DataFrame]:
    groups = [
        ["CUS HORTICULTURE", "CUS411S", "Windhoek Main Campus", "Full-time", "Yes"],
        ["CUS GROUP A", "CUS411S", "Windhoek Main Campus", "Full-time", "Yes"],
        ["CUS GROUP B", "CUS411S", "Windhoek Main Campus", "Full-time", "Yes"],
        ["ICT GROUP A", "ICT521S", "Windhoek Main Campus", "Full-time", "Yes"],
        ["ICT GROUP B", "ICT521S", "Windhoek Main Campus", "Full-time", "Yes"],
        ["ICT BOA-EENANHA", "ICT521S", "Eenhana Satellite Campus", "Full-time", "Yes"],
        ["ICT GREY", "ICT521S", "Windhoek Main Campus", "Full-time", "Yes"],
        ["ICT Distance", "ICT521S", "Distance / Online", "Part-time", "Yes"],
    ]
    students = []
    enrolments = []
    student_number = 910000001
    for group_name, course_code, *_ in groups:
        for _ in range(12):
            index = student_number - 910000000
            students.append(
                [
                    str(student_number),
                    f"PilotSurname{index:03d}",
                    f"P{index:03d}",
                    f"Pilot Student {index:03d}",
                    "Yes",
                ]
            )
            enrolments.append([str(student_number), group_name, course_code, "Yes"])
            student_number += 1

    return {
        "Lecturers": pd.DataFrame(
            [
                [
                    "1008977",
                    "Ms",
                    "Maria Matias",
                    "Master of Science in Information Technology",
                    "PLACEHOLDER-ID-1008977",
                    "PLACEHOLDER-PAYE-1008977",
                    "PLACEHOLDER ADDRESS",
                    "0810000000",
                    460,
                    "Windhoek Main Campus",
                    "2026-02-09",
                    "2026-06-05",
                    "Yes",
                ]
            ],
            columns=SHEET_COLUMNS["Lecturers"],
        ),
        "Courses": pd.DataFrame(
            [
                ["CUS411S", "Computer User Skills", "Computing and Informatics", "Informatics", "0183-0102", "Yes"],
                ["ICT521S", "Information Competence", "Computing and Informatics", "Informatics", "0183-0102", "Yes"],
            ],
            columns=SHEET_COLUMNS["Courses"],
        ),
        "Groups": pd.DataFrame(groups, columns=SHEET_COLUMNS["Groups"]),
        "Students": pd.DataFrame(students, columns=SHEET_COLUMNS["Students"]),
        "Group_Enrolments": pd.DataFrame(enrolments, columns=SHEET_COLUMNS["Group_Enrolments"]),
        "Timetable": pd.DataFrame(TIMETABLE_ROWS, columns=SHEET_COLUMNS["Timetable"]),
        "Academic_Calendar": pd.DataFrame(CALENDAR_ROWS, columns=SHEET_COLUMNS["Academic_Calendar"]),
    }


def _format_sheet(writer: pd.ExcelWriter, sheet_name: str, fill_color: str) -> None:
    worksheet = writer.sheets[sheet_name]
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor=fill_color)
    _autosize_columns(writer, sheet_name)


def create_maria_pilot_workbook(output_path: str | Path = PILOT_PATH) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _pilot_instructions_df().to_excel(writer, sheet_name="Instructions", index=False)
        _data_dictionary_df().to_excel(writer, sheet_name="Data_Dictionary", index=False)
        _allowed_values_df().to_excel(writer, sheet_name="Allowed_Values", index=False)
        for sheet_name in META_SHEETS:
            _format_sheet(writer, sheet_name, "E2F0D9")
        for sheet_name, df in _pilot_dataframes().items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            _format_sheet(writer, sheet_name, "D9EAF7")
    return output


def main() -> None:
    path = create_maria_pilot_workbook()
    print("Maria Matias April 2026 pilot workbook created.")
    print("PILOT WARNING: reconstructed from a submitted claim, not an official approved timetable.")
    print("Sensitive values are placeholders and students are dummy-only.")
    print(f"Workbook path: {path}")
    print(f"Expected April 2026 totals: {EXPECTED_SESSIONS} sessions, {EXPECTED_HOURS:.2f} hours, {EXPECTED_AMOUNT:.2f} amount.")


if __name__ == "__main__":
    main()
